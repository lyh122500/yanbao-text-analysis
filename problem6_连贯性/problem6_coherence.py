#!/usr/bin/env python3
"""Deterministic, corpus-calibrated coherence proxies for Chinese analyst reports."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import jieba
import joblib
import numpy as np
import openpyxl
import scipy
import sklearn
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.decomposition import PCA
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.preprocessing import StandardScaler

HERE = Path(__file__).resolve().parent
VERSION = "1.0.0"
REFERENCE_VERSION = "problem6-reference-v1"
HAN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
SENTENCE_BOUNDARY = re.compile(r"<\?>|[。！？!?]+")
PLACEHOLDERS = {"", "无", "暂无", "无内容", "无有效内容", "nan", "none", "null"}

# 固定表，不从本批数据反推，避免每次运行改变定义。
STOPWORDS = {
    "我们", "公司", "认为", "预计", "目前", "已经", "以及", "其中", "对于", "由于", "因此",
    "可以", "可能", "主要", "相关", "进行", "通过", "随着", "同比", "环比", "分别", "未来",
    "市场", "行业", "业务", "产品", "实现", "增长", "收入", "利润", "方面", "情况", "较为",
    "一个", "进一步", "同时", "将会", "仍然", "保持", "有所", "报告", "分析师", "研究员",
}
CONNECTIVES = (
    "因此", "所以", "因而", "由此", "从而", "于是", "故而", "可见", "综上", "总之",
    "但是", "然而", "不过", "尽管", "虽然", "相反", "相比之下", "另一方面", "与此同时",
    "此外", "另外", "再者", "同时", "进一步", "其次", "最后", "首先", "具体来看",
    "换言之", "也就是说", "例如", "比如", "尤其是", "事实上", "值得注意的是",
)
REFERENTIALS = (
    "这", "该", "其", "此", "上述", "前述", "其中", "对此", "由此", "这些", "那些",
    "前者", "后者", "二者", "两者", "这一", "这种", "此类", "该类", "它", "他们",
)
FEATURES = [
    "语义局部超额", "词汇局部超额", "相邻共享实词比例", "主题聚焦度",
    "显式连接覆盖率", "指代承接覆盖率", "非零语义转移比例",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def dump_json(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def normalize_text(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u00ad", "").replace("\u200b", "").replace("\ufeff", "")
    text = re.sub(r"[\t\r\n ]+", "", text)
    return text


def split_report(value):
    """Return unique cleaned sentences, section ids, and exact duplicates removed."""
    text = normalize_text(value)
    if text.lower() in PLACEHOLDERS or not HAN.search(text):
        return [], [], 0
    sentences, sections, seen, duplicate_count = [], [], set(), 0

    def append_piece(piece, section_id):
        nonlocal duplicate_count
        if not HAN.search(piece):
            return
        if piece in seen:
            duplicate_count += 1
            return
        seen.add(piece)
        sentences.append(piece)
        sections.append(section_id)

    section_id = 1
    start = 0
    for match in SENTENCE_BOUNDARY.finditer(text):
        piece = text[start:match.start()].strip("；;，,：:")
        append_piece(piece, section_id)
        if match.group() == "<?>":
            section_id += 1
        start = match.end()
    piece = text[start:].strip("；;，,：:")
    append_piece(piece, section_id)
    return sentences, sections, duplicate_count


def count_han(value) -> int:
    return len(HAN.findall(normalize_text(value)))


def content_tokens(sentence: str, tokenizer=None):
    tokenizer = tokenizer or jieba.Tokenizer()
    terms = set()
    for token in tokenizer.cut(sentence, HMM=False):
        token = token.strip()
        if len(HAN.findall(token)) >= 2 and token not in STOPWORDS:
            terms.add(token)
    return terms


def cosine_matrix(x):
    # TfidfVectorizer produces L2-normalized rows; product is cosine similarity.
    return (x @ x.T).toarray().astype(float, copy=False)


def pair_values(matrix, token_sets):
    n = len(token_sets)
    if n < 2:
        return np.array([]), np.array([]), np.array([]), np.array([])
    adjacent_semantic = np.diag(matrix, 1)
    nonadjacent_semantic = matrix[np.triu_indices(n, 2)] if n >= 3 else np.array([])
    adjacent_lexical = []
    nonadjacent_lexical = []
    for i in range(n):
        for j in range(i + 1, n):
            union = token_sets[i] | token_sets[j]
            score = len(token_sets[i] & token_sets[j]) / len(union) if union else 0.0
            (adjacent_lexical if j == i + 1 else nonadjacent_lexical).append(score)
    return (np.asarray(adjacent_semantic), np.asarray(nonadjacent_semantic),
            np.asarray(adjacent_lexical), np.asarray(nonadjacent_lexical))


def mean_or_none(values):
    return float(np.mean(values)) if len(values) else None


def starts_with(sentence, prefixes) -> bool:
    sentence = sentence.lstrip("（(【[“\"'《")
    return sentence.startswith(prefixes)


def analyze_document(sentences, sections, matrix, token_sets):
    n = len(sentences)
    a_sem, na_sem, a_lex, na_lex = pair_values(matrix, token_sets)
    a_sem_mean, na_sem_mean = mean_or_none(a_sem), mean_or_none(na_sem)
    a_lex_mean, na_lex_mean = mean_or_none(a_lex), mean_or_none(na_lex)
    centroid = np.asarray(matrix.mean(axis=1)).ravel() if n else np.array([])
    # cos(sentence, normalized mean TF-IDF vector) can be derived from the Gram matrix.
    centroid_norm = float(np.sqrt(matrix.mean())) if n and matrix.mean() > 0 else 0.0
    topic_focus = float(np.mean(centroid / centroid_norm)) if centroid_norm else 0.0
    within = [a_sem[i] for i in range(max(0, n - 1)) if sections[i] == sections[i + 1]]
    across = [a_sem[i] for i in range(max(0, n - 1)) if sections[i] != sections[i + 1]]
    shared = [bool(token_sets[i] & token_sets[i + 1]) for i in range(max(0, n - 1))]
    transition_n = max(0, n - 1)
    metrics = {
        "句数": n,
        "段数": len(set(sections)),
        "相邻转移数": transition_n,
        "相邻句语义相似度": a_sem_mean,
        "非相邻句语义相似度": na_sem_mean,
        "语义局部超额": (a_sem_mean - na_sem_mean) if a_sem_mean is not None and na_sem_mean is not None else None,
        "零语义转移比例": float(np.mean(a_sem <= 1e-12)) if len(a_sem) else None,
        "非零语义转移比例": float(np.mean(a_sem > 1e-12)) if len(a_sem) else None,
        "相邻句词汇Jaccard": a_lex_mean,
        "非相邻句词汇Jaccard": na_lex_mean,
        "词汇局部超额": (a_lex_mean - na_lex_mean) if a_lex_mean is not None and na_lex_mean is not None else None,
        "相邻共享实词比例": float(np.mean(shared)) if shared else None,
        "主题聚焦度": topic_focus if n else None,
        "首末句语义相似度": float(matrix[0, -1]) if n >= 2 else None,
        "段内相邻语义相似度": mean_or_none(within),
        "跨段相邻语义相似度": mean_or_none(across),
        "显式连接覆盖率": sum(starts_with(s, CONNECTIVES) for s in sentences[1:]) / transition_n if transition_n else None,
        "指代承接覆盖率": sum(starts_with(s, REFERENTIALS) for s in sentences[1:]) / transition_n if transition_n else None,
    }
    # Freeze numeric precision before corpus ranking. fit_transform() and transform()
    # can differ by a few float32 ULPs; without quantization those harmless changes
    # may swap percentile ranks and make a frozen-reference replay look different.
    metrics = {key: round(value, 7) if isinstance(value, float) else value
               for key, value in metrics.items()}
    return metrics, a_sem, a_lex


def percentile(values, sorted_reference):
    reference = np.asarray(sorted_reference, dtype=float)
    if not len(reference):
        return np.full(len(values), np.nan)
    left = np.searchsorted(reference, values, side="left")
    right = np.searchsorted(reference, values, side="right")
    if len(reference) == 1:
        return np.full(len(values), 50.0)
    return 100.0 * ((left + right - 1) / 2.0) / (len(reference) - 1)


def fit_or_apply_scores(documents, reference=None):
    eligible = [i for i, d in enumerate(documents) if d["计算状态"] == "可比较"]
    if not eligible:
        raise ValueError("没有满足至少100汉字且至少5句的研报，无法建立参考分布")
    x = np.asarray([[documents[i][f] for f in FEATURES] for i in eligible], dtype=float)
    if reference is None:
        feature_reference = {f: np.sort(x[:, j]) for j, f in enumerate(FEATURES)}
        scaler = StandardScaler().fit(x)
        pca = PCA(n_components=1, svd_solver="full").fit(scaler.transform(x))
        pc = pca.transform(scaler.transform(x)).ravel()
        percent_matrix = np.column_stack([percentile(x[:, j], feature_reference[f]) for j, f in enumerate(FEATURES)])
        equal = percent_matrix.mean(axis=1)
        direction = 1.0 if np.corrcoef(pc, equal)[0, 1] >= 0 else -1.0
        pc *= direction
        pc_reference = np.sort(pc)
        reference = {
            "reference_version": REFERENCE_VERSION,
            "program_version": VERSION,
            "features": FEATURES,
            "feature_reference": feature_reference,
            "scaler": scaler,
            "pca": pca,
            "pca_direction": direction,
            "pca_reference": pc_reference,
            "eligible_count": len(eligible),
        }
    else:
        if reference.get("reference_version") != REFERENCE_VERSION or reference.get("features") != FEATURES:
            raise ValueError("参考模型版本或特征定义不兼容")
        feature_reference = reference["feature_reference"]
        scaler, pca, direction = reference["scaler"], reference["pca"], reference["pca_direction"]
        percent_matrix = np.column_stack([percentile(x[:, j], feature_reference[f]) for j, f in enumerate(FEATURES)])
        equal = percent_matrix.mean(axis=1)
        pc = pca.transform(scaler.transform(x)).ravel() * direction
        pc_reference = reference["pca_reference"]
    pca_score = percentile(pc, pc_reference)
    for row_idx, eq, raw_pc, score in zip(eligible, equal, pc, pca_score):
        documents[row_idx]["等权连贯分"] = float(eq)
        documents[row_idx]["PCA原始分"] = float(raw_pc)
        documents[row_idx]["PCA连贯分"] = float(score)
    return reference


def deterministic_shuffle_validation(documents, doc_data, max_docs=500, repeats=10):
    eligible = [i for i, d in enumerate(documents) if d["计算状态"] == "可比较"]
    chosen = sorted(eligible, key=lambda i: hashlib.sha256(str(documents[i]["源Excel行号"]).encode()).hexdigest())[:max_docs]
    comparisons = []
    for i in chosen:
        sentences, sections, matrix, token_sets = doc_data[i]
        original = documents[i]
        rng = np.random.default_rng(documents[i]["源Excel行号"])
        shuffled = []
        for _ in range(repeats):
            order = rng.permutation(len(sentences))
            m = matrix[np.ix_(order, order)]
            t = [token_sets[j] for j in order]
            s = [sections[j] for j in order]
            sm, _, _ = analyze_document([sentences[j] for j in order], s, m, t)
            shuffled.append(sm)
        comparisons.append({
            "source_row": documents[i]["源Excel行号"],
            "semantic_original": original["相邻句语义相似度"],
            "semantic_shuffled": float(np.mean([d["相邻句语义相似度"] for d in shuffled])),
            "lexical_original": original["相邻句词汇Jaccard"],
            "lexical_shuffled": float(np.mean([d["相邻句词汇Jaccard"] for d in shuffled])),
            "semantic_excess_original": original["语义局部超额"],
            "semantic_excess_shuffled": float(np.mean([d["语义局部超额"] for d in shuffled])),
            "lexical_excess_original": original["词汇局部超额"],
            "lexical_excess_shuffled": float(np.mean([d["词汇局部超额"] for d in shuffled])),
        })
    def summarize(a, b):
        av = np.asarray([r[a] for r in comparisons])
        bv = np.asarray([r[b] for r in comparisons])
        return {"原文均值": float(av.mean()), "打乱均值": float(bv.mean()),
                "原文减打乱": float((av-bv).mean()), "原文高于打乱比例": float(np.mean(av > bv))}
    return {
        "method": f"按源行号哈希固定抽取{len(chosen)}篇，每篇以源行号为随机种子打乱{repeats}次",
        "sample_size": len(chosen), "repeats": repeats,
        "相邻句语义相似度": summarize("semantic_original", "semantic_shuffled"),
        "相邻句词汇Jaccard": summarize("lexical_original", "lexical_shuffled"),
        "语义局部超额": summarize("semantic_excess_original", "semantic_excess_shuffled"),
        "词汇局部超额": summarize("lexical_excess_original", "lexical_excess_shuffled"),
        "sample_rows": comparisons,
    }


def create_review_sample(path, documents, raw_by_index, n_each=20):
    eligible = [i for i, d in enumerate(documents) if d.get("PCA连贯分") is not None]
    groups = {
        "PCA低分": sorted(eligible, key=lambda i: documents[i]["PCA连贯分"])[:n_each],
        "PCA高分": sorted(eligible, key=lambda i: -documents[i]["PCA连贯分"])[:n_each],
        "两种综合分分歧大": sorted(eligible, key=lambda i: -abs(documents[i]["PCA连贯分"]-documents[i]["等权连贯分"]))[:n_each],
        "中位附近随机": sorted(eligible, key=lambda i: (abs(documents[i]["PCA连贯分"]-50), hashlib.sha256(str(i).encode()).hexdigest()))[:n_each],
    }
    rows, seen = [], set()
    for group, indices in groups.items():
        for i in indices:
            key = (group, i)
            if key in seen:
                continue
            seen.add(key)
            d = documents[i]
            text = normalize_text(raw_by_index[i])
            rows.append({"抽样层": group, "源Excel行号": d["源Excel行号"], "公司简称": d.get("公司简称"),
                         "PCA连贯分": d.get("PCA连贯分"), "等权连贯分": d.get("等权连贯分"),
                         "相邻句语义相似度": d.get("相邻句语义相似度"), "语义局部超额": d.get("语义局部超额"),
                         "相邻共享实词比例": d.get("相邻共享实词比例"), "文本开头_仅供核验": text[:500],
                         "人工整体连贯_1至5": "", "人工备注": "", "标注人": ""})
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def write_workbook(path, documents, metadata_names):
    metric_names = ["计算状态", "有效汉字数", "重复句数_已剔除", "句数", "段数", "相邻转移数",
        "相邻句语义相似度", "非相邻句语义相似度", "语义局部超额", "零语义转移比例",
        "相邻句词汇Jaccard", "非相邻句词汇Jaccard", "词汇局部超额", "相邻共享实词比例",
        "主题聚焦度", "首末句语义相似度", "段内相邻语义相似度", "跨段相邻语义相似度",
        "显式连接覆盖率", "指代承接覆盖率", "等权连贯分", "PCA原始分", "PCA连贯分"]
    headers = ["源Excel行号"] + metadata_names + metric_names
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "连贯性指标"
    ws.append(headers)
    for d in documents:
        ws.append([d.get(h) for h in headers])
    ws.freeze_panes = f"{get_column_letter(len(metadata_names)+2)}2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 62
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = PatternFill("solid", fgColor="29445C")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 21 if title in metric_names else 15
        for c in ws.iter_cols(min_col=col, max_col=col, min_row=2, max_row=ws.max_row):
            v = c[0]
            v.font = Font(name="Arial", size=9)
            if title == "fordate":
                v.number_format = "yyyy-mm-dd"
            elif title == "stkcd":
                v.number_format = "000000"
            elif title.endswith("比例") or title.endswith("覆盖率"):
                v.number_format = "0.00%"
            elif title in metric_names and title not in {"计算状态", "有效汉字数", "重复句数_已剔除", "句数", "段数", "相邻转移数"}:
                v.number_format = "0.0000"
    ws.auto_filter.ref = ws.dimensions
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=HERE.parent / "综合分析结果.xlsx")
    parser.add_argument("--output-dir", type=Path, default=HERE / "output")
    parser.add_argument("--reference", type=Path, help="冻结参考模型；省略时用本批数据拟合并保存")
    parser.add_argument("--max-features", type=int, default=50000)
    parser.add_argument("--shuffle-docs", type=int, default=500)
    parser.add_argument("--shuffle-repeats", type=int, default=10)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    headings = list(next(rows))
    if "研报文本" not in headings:
        raise KeyError("源文件缺少‘研报文本’列")
    metadata_names = [h for h in ["fordate", "序号", "stkcd", "公司简称", "证券公司"] if h in headings]
    records, all_sentences, slices, raw_by_index = [], [], [], []
    for source_row, values in enumerate(rows, 2):
        record = dict(zip(headings, values))
        sentences, sections, duplicate_count = split_report(record.get("研报文本"))
        start = len(all_sentences); all_sentences.extend(sentences)
        slices.append((start, len(all_sentences), sentences, sections, duplicate_count))
        records.append({"源Excel行号": source_row, **{k: record.get(k) for k in metadata_names},
                        "有效汉字数": count_han(record.get("研报文本"))})
        raw_by_index.append(record.get("研报文本"))
    wb.close()

    loaded_reference = joblib.load(args.reference) if args.reference else None
    vectorizer = loaded_reference["vectorizer"] if loaded_reference else TfidfVectorizer(
        analyzer="char", ngram_range=(2, 4), min_df=5, max_features=args.max_features,
        sublinear_tf=True, norm="l2", lowercase=False, dtype=np.float32)
    # Always use the same transform path. sklearn's sparse fit_transform and
    # transform paths may differ in the last float32 bits, which can alter ties
    # in a percentile score even though the underlying metric is unchanged.
    if loaded_reference:
        x_all = vectorizer.transform(all_sentences)
    else:
        vectorizer.fit(all_sentences)
        x_all = vectorizer.transform(all_sentences)
    tokenizer = jieba.Tokenizer()
    documents, doc_data = [], []
    evidence_path = args.output_dir / "句际转移证据.jsonl"
    with evidence_path.open("w", encoding="utf-8") as evidence:
        for i, (start, end, sentences, sections, duplicate_count) in enumerate(slices):
            x = x_all[start:end]
            matrix = cosine_matrix(x)
            token_sets = [content_tokens(s, tokenizer) for s in sentences]
            metrics, adjacent_semantic, adjacent_lexical = analyze_document(sentences, sections, matrix, token_sets)
            han = records[i]["有效汉字数"]
            if not sentences:
                status = "空白或占位文本"
            elif han < 100 or len(sentences) < 5:
                status = "短文本_不进入参考分布"
            else:
                status = "可比较"
            d = {**records[i], "计算状态": status, "重复句数_已剔除": duplicate_count, **metrics,
                 "等权连贯分": None, "PCA原始分": None, "PCA连贯分": None}
            documents.append(d); doc_data.append((sentences, sections, matrix, token_sets))
            for j in range(max(0, len(sentences)-1)):
                evidence.write(json.dumps({"source_row": d["源Excel行号"], "left_sentence_id": j+1,
                    "right_sentence_id": j+2, "section_boundary": sections[j] != sections[j+1],
                    "semantic_cosine": float(adjacent_semantic[j]), "lexical_jaccard": float(adjacent_lexical[j]),
                    "shared_content_tokens": sorted(token_sets[j] & token_sets[j+1])[:20]},
                    ensure_ascii=False, sort_keys=True) + "\n")

    score_reference = loaded_reference and loaded_reference["score_reference"]
    score_reference = fit_or_apply_scores(documents, score_reference)
    if not loaded_reference:
        reference_path = args.output_dir / "reference.joblib"
        joblib.dump({"vectorizer": vectorizer, "score_reference": score_reference,
                     "source_sha256": sha256(args.input), "created_utc": datetime.now(timezone.utc).isoformat(),
                     "versions": {"python": platform.python_version(), "numpy": np.__version__,
                                  "scipy": scipy.__version__, "sklearn": sklearn.__version__,
                                  "jieba": getattr(jieba, "__version__", "unknown")}}, reference_path, compress=3)
    else:
        reference_path = args.reference

    validation = deterministic_shuffle_validation(documents, doc_data, args.shuffle_docs, args.shuffle_repeats)
    validation_rows = validation.pop("sample_rows")
    dump_json(args.output_dir / "shuffle_validation.json", validation)
    with (args.output_dir / "shuffle_validation_rows.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(validation_rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(validation_rows)
    create_review_sample(args.output_dir / "连贯性人工核验样本.csv", documents, raw_by_index)
    output_xlsx = args.output_dir / "问题6_连贯性结果.xlsx"
    write_workbook(output_xlsx, documents, metadata_names)

    statuses = Counter(d["计算状态"] for d in documents)
    eligible = [d for d in documents if d["计算状态"] == "可比较"]
    summary = {
        "program_version": VERSION, "source_rows": len(documents), "sentence_count": len(all_sentences),
        "exact_duplicate_sentences_removed": sum(d["重复句数_已剔除"] for d in documents),
        "status_counts": dict(statuses), "reference_eligible": len(eligible),
        "PCA_explained_variance_ratio": float(score_reference["pca"].explained_variance_ratio_[0]),
        "PCA_loadings_after_orientation": {f: float(v * score_reference["pca_direction"])
                                            for f, v in zip(FEATURES, score_reference["pca"].components_[0])},
        "metric_medians": {k: float(np.median([d[k] for d in eligible])) for k in FEATURES + ["PCA连贯分", "等权连贯分"]},
        "shuffle_validation": validation,
    }
    dump_json(args.output_dir / "summary.json", summary)
    artifacts = [output_xlsx, args.output_dir / "summary.json", args.output_dir / "shuffle_validation.json"]
    if not args.reference:
        artifacts.append(reference_path)
    manifest = {"program_version": VERSION, "source": str(args.input), "source_sha256": sha256(args.input),
                "reference": str(reference_path), "reference_sha256": sha256(reference_path),
                "outputs": {p.name: sha256(p) for p in artifacts}}
    dump_json(args.output_dir / "manifest.json", manifest)
    print(json.dumps({"rows": len(documents), "sentences": len(all_sentences), "status": dict(statuses),
                      "xlsx": str(output_xlsx)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
