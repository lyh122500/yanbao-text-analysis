# -*- coding: utf-8 -*-
"""Standalone, deterministic surface-complexity/readability proxies. See README."""
import argparse
import bisect
import hashlib
import json
import math
import platform
import re
import statistics
import unicodedata
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
VERSION = '1.0.0'
HAN = re.compile(r'[\u3400-\u4dbf\u4e00-\u9fff\U00020000-\U0002ebef]')
NUMBER = re.compile(r'(?<![A-Za-z])[-+]?\d+(?:,\d{3})*(?:\.\d+)?%?')
FEATURES = ('句均汉字数', '句均分段数', '分段均汉字数', '非一级常用字比例')
HEADINGS = {'财报点评', '公司研究', '公司点评', '投资要点', '投资摘要', '风险提示', '风险因素', '公司简评'}
INVALID = {'无', '暂无', '无内容', '无有效内容', 'null', 'none', 'nan', 'n/a'}
METRICS = ['状态','质量提示','复杂性_代理','可读性_代理','阅读负担_代理','汉字数','句数','标点分段数',
           *FEATURES, '一级常用字数','次常用字数','表外字数','字表统计汉字数','次常用字比例','表外字比例',
           '数字串数','每千汉字数字串','占位分隔符数','最长句汉字数','长句比例_超80字','移除标题行数',
           '非一级字命中','最长句证据','复杂性_分号分句对照']

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def load_chars(path, expected):
    chars = []
    for line in Path(path).read_text(encoding='utf-8-sig').splitlines():
        if line.strip():
            c = line.split()[0]
            if len(c) != 1 or not HAN.fullmatch(c):
                raise ValueError('字表首列必须是单个汉字: ' + line)
            chars.append(c)
    if len(chars) != expected or len(set(chars)) != expected:
        raise ValueError('字表数量/唯一性不符: ' + str(path))
    return set(chars)

def clean(text):
    # NFKC makes full-width decimal numerals/punctuation consistent.
    text = unicodedata.normalize('NFKC', str(text))
    text = re.sub('[\u200b-\u200d\ufeff]', '', text)
    text = re.sub('[\xad‐‑‒–—−]', '-', text)
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    lines, removed = [], 0
    for line in text.split('\n'):
        if line.strip().rstrip(':：') in HEADINGS:
            removed += 1
        else:
            lines.append(line)
    text = '\n'.join(lines)
    # Explicit separators and blank lines are boundaries; single wraps are not.
    text = re.sub(r'<\?>|\n\s*\n', '\u2029', text)
    text = re.sub(r'(?<=[\u4e00-\u9fff])\n(?=[\u4e00-\u9fff])', '', text)
    text = text.replace('\n', ' ')
    return text, removed

def split_sentences(text, semicolon=False):
    boundary = r'[。！？!?\u2029]+' if not semicolon else r'[。！？!?;；\u2029]+'
    # An ASCII full stop is a boundary only next to Han/closing quote, never 0.64 or A股.
    text = re.sub(r'(?<=[\u4e00-\u9fff”」])\.(?!\d)', '。', text)
    return [s.strip() for s in re.split(boundary, text) if HAN.search(s)]

def split_clauses(sentence):
    # Protect thousands separators before comma-based segmentation.
    protected = re.sub(r'(?<=\d),(?=\d{3}(?:\D|$))', '\u2060', sentence)
    return [s.replace('\u2060', ',').strip() for s in re.split(r'[,，;；:：]+', protected) if HAN.search(s)]

def extract(raw, company, primary, secondary, semicolon=False):
    out = {k: None for k in METRICS}
    if raw is None or not str(raw).strip():
        out.update(状态='文本为空', 质量提示='不计分')
        return out
    if str(raw).strip().lower() in INVALID:
        out.update(状态='占位文本', 质量提示='不计分')
        return out
    text, removed = clean(raw)
    sentences = split_sentences(text, semicolon)
    counts = [len(HAN.findall(s)) for s in sentences]
    clauses = [c for s in sentences for c in split_clauses(s)]
    chars = HAN.findall(text)
    # Remove exact company mentions from familiarity only, preserving structural lengths.
    lexical = text
    if company:
        company = unicodedata.normalize('NFKC', str(company)).strip()
        if company:
            lexical = lexical.replace(company, '')
    lexical_chars = HAN.findall(lexical)
    n = len(lexical_chars)
    first = sum(c in primary for c in lexical_chars)
    second = sum(c in secondary for c in lexical_chars)
    outside = n - first - second
    flags = []
    if not sentences or len(chars) < 100 or len(sentences) < 3:
        flags.append('短文本不足100汉字或3句')
    if counts and max(counts) > 500:
        flags.append('超长句_检查抽取或缺标点')
    if len(re.findall(r'\.{4,}', text)) >= 3:
        flags.append('疑似目录残片')
    if n == 0:
        flags.append('字表分母为零')
    out.update({'状态': '可计算' if not flags else ('短文本' if flags[0].startswith('短文本') else '需复核'),
        '质量提示': ';'.join(flags), '汉字数':len(chars), '句数':len(sentences),
        '标点分段数':len(clauses), '句均汉字数': statistics.mean(counts) if counts else None,
        '句均分段数':len(clauses)/len(sentences) if sentences else None,
        '分段均汉字数':sum(len(HAN.findall(c)) for c in clauses)/len(clauses) if clauses else None,
        '一级常用字数':first,'次常用字数':second,'表外字数':outside,'字表统计汉字数':n,
        '非一级常用字比例':(second+outside)/n if n else None,
        '次常用字比例':second/n if n else None,'表外字比例':outside/n if n else None,
        '数字串数':len(NUMBER.findall(text)),
        '每千汉字数字串':1000*len(NUMBER.findall(text))/len(chars) if chars else None,
        '占位分隔符数':str(raw).count('<?>'), '最长句汉字数':max(counts) if counts else None,
        '长句比例_超80字':sum(x>80 for x in counts)/len(counts) if counts else None,
        '移除标题行数':removed,
        '非一级字命中': ';'.join(f'{c}:{v}' for c,v in sorted(Counter(c for c in lexical_chars if c not in primary).items(),key=lambda x:(-x[1],x[0]))[:15]),
        '最长句证据': sentences[counts.index(max(counts))][:500] if counts else ''})
    return out

def distribution(values):
    count = Counter(values)
    return {'values':sorted(count), 'counts':[count[x] for x in sorted(count)]}

def build_reference(metrics, dictionary_hashes, input_hash):
    good = [m for m in metrics if m['状态']=='可计算']
    if len(good) < 2:
        raise ValueError('至少需要两条可计算文本作为参考分布')
    return {'version':VERSION,'input_sha256':input_hash,'dictionary_sha256':dictionary_hashes,
            'reference_n':len(good),'features':{k:distribution([m[k] for m in good]) for k in FEATURES}}

def prepare_reference(ref):
    prepared = {}
    for k in FEATURES:
        d = ref['features'][k]
        cumulative=[0]
        for c in d['counts']: cumulative.append(cumulative[-1]+c)
        prepared[k]=(d['values'], cumulative)
    return prepared

def rank(value, dist):
    values,cumulative=dist
    lo=bisect.bisect_left(values,value);hi=bisect.bisect_right(values,value)
    return (cumulative[lo]+.5*(cumulative[hi]-cumulative[lo]))/cumulative[-1]*100

def score(m, ref):
    if m['状态']!='可计算':
        return m
    ranks={k:rank(m[k],ref[k]) for k in FEATURES}
    m['复杂性_代理']=(ranks[FEATURES[0]]+ranks[FEATURES[1]])/2
    m['阅读负担_代理']=(ranks[FEATURES[2]]+ranks[FEATURES[3]])/2
    m['可读性_代理']=100-m['阅读负担_代理']
    return m

def corr(a,b):
    if len(a)<2:return None
    ma=statistics.mean(a);mb=statistics.mean(b)
    den=math.sqrt(sum((x-ma)**2 for x in a)*sum((x-mb)**2 for x in b))
    return sum((x-ma)*(y-mb) for x,y in zip(a,b))/den if den else None

def summarize(rows):
    summary={'rows':len(rows),'status':dict(Counter(m['状态'] for m in rows)), 'statistics':{}}
    for key in ['复杂性_代理','可读性_代理',*FEATURES]:
        x=sorted(m[key] for m in rows if m[key] is not None)
        summary['statistics'][key]={'n':len(x),'mean':statistics.mean(x),'median':statistics.median(x),'min':min(x),'max':max(x)} if x else {}
    good=[m for m in rows if m['复杂性_代理'] is not None]
    summary['pearson']={
        '复杂性_vs_可读性':corr([m['复杂性_代理'] for m in good],[m['可读性_代理'] for m in good]),
        '复杂性_vs_汉字数':corr([m['复杂性_代理'] for m in good],[m['汉字数'] for m in good]),
        '可读性_vs_汉字数':corr([m['可读性_代理'] for m in good],[m['汉字数'] for m in good])}
    return summary

def save_workbook(path, records, metrics):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='复杂性可读性'
    headers=['源Excel行号','fordate','序号','stkcd','公司简称','证券公司']+METRICS
    ws.append(headers)
    for record,m in zip(records,metrics):
        ws.append(record+[m[k] for k in METRICS])
    ws.freeze_panes='G2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234E70');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=42
    for i,key in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(i)].width=42 if key in ['质量提示','非一级字命中','最长句证据','复杂性_分号分句对照'] else 19
    for row in ws.iter_rows(min_row=2):
        for cell,key in zip(row,headers):
            cell.font=Font(name='Arial',size=10)
            if key=='fordate':cell.number_format='yyyy-mm-dd'
            elif isinstance(cell.value,float):cell.number_format='0.00%' if '比例' in key else '0.0000'
            if isinstance(cell.value,str):cell.data_type='s'  # never interpret source text as Excel formulas
    wb.save(path)

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',type=Path,default=ROOT/'综合分析结果.xlsx')
    p.add_argument('--output-dir',type=Path,default=HERE/'output')
    p.add_argument('--reference',type=Path,help='冻结参考分布；默认拟合当前输入并保存')
    args=p.parse_args()
    paths=[HERE/'dict/现代汉语常用字表_一级常用字2500.txt',HERE/'dict/现代汉语常用字表_次常用字1000.txt']
    primary=load_chars(paths[0],2500);secondary=load_chars(paths[1],1000)
    if primary & secondary:raise ValueError('两级字表重叠')
    hashes={x.name:sha(x) for x in paths}
    wb=openpyxl.load_workbook(args.input,read_only=True,data_only=True);ws=wb['Sheet1']
    it=ws.iter_rows(values_only=True);header=next(it);idx={k:i for i,k in enumerate(header)}
    records=[];metrics=[];alternatives=[]
    for rn,row in enumerate(it,2):
        get=lambda key:row[idx[key]]
        records.append([rn]+[get(k) for k in ['fordate','序号','stkcd','公司简称','证券公司']])
        metrics.append(extract(get('研报文本'),get('公司简称'),primary,secondary))
        alternatives.append(extract(get('研报文本'),get('公司简称'),primary,secondary,True))
    wb.close();input_hash=sha(args.input)
    ref=json.loads(args.reference.read_text()) if args.reference else build_reference(metrics,hashes,input_hash)
    if ref['version']!=VERSION or ref['dictionary_sha256']!=hashes:raise ValueError('参考分布版本或字表哈希不一致')
    prepared=prepare_reference(ref)
    metrics=[score(m,prepared) for m in metrics]
    alternatives=[score(m,prepared) for m in alternatives]
    for m, alt in zip(metrics, alternatives):
        m['复杂性_分号分句对照'] = alt['复杂性_代理']
    summary=summarize(metrics)
    paired=[(a,b) for a,b in zip(metrics,alternatives) if a['复杂性_代理'] is not None and b['复杂性_代理'] is not None]
    summary['semicolon_sensitivity']={'paired_n':len(paired),'complexity_pearson':corr([a['复杂性_代理'] for a,b in paired],[b['复杂性_代理'] for a,b in paired]),'definition':'仅将分号改作句末，沿用同一冻结参考分布'}
    args.output_dir.mkdir(parents=True,exist_ok=True)
    save_workbook(args.output_dir/'问题3_复杂性可读性结果.xlsx',records,metrics)
    payload=[dict(zip(['源Excel行号','fordate','序号','stkcd','公司简称','证券公司'],r),**m) for r,m in zip(records,metrics)]
    canonical=json.dumps(payload,ensure_ascii=False,sort_keys=True,default=str,allow_nan=False)
    (args.output_dir/'结果明细.json').write_text(canonical,encoding='utf-8')
    (args.output_dir/'reference.json').write_text(json.dumps(ref,ensure_ascii=False,indent=2),encoding='utf-8')
    (args.output_dir/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    manifest={'method_version':VERSION,'input_sha256':input_hash,'script_sha256':sha(__file__),
              'dictionary_sha256':hashes,'python':platform.python_version(),'openpyxl':openpyxl.__version__,
              'reference_sha256':sha(args.output_dir/'reference.json'),'results_sha256':hashlib.sha256(canonical.encode()).hexdigest(),
              'reference_mode':'reuse' if args.reference else 'fit', 'interpretation':'未经读者理解测试验证的相对代理指标，非教育年限/文献原式'}
    (args.output_dir/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
