"""Audit source offsets, aggregation, workbook values and optional exact rerun."""
import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl
from problem5_extremity import HERE, LABELS, dump, normalized, sha


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--input", type=Path, default=HERE.parent / "综合分析结果.xlsx")
    p.add_argument("--output-dir", type=Path, default=HERE / "output")
    p.add_argument("--replay-dir", type=Path)
    args = p.parse_args()
    manifest = json.loads((args.output_dir / "manifest.json").read_text())
    assert sha(args.input) == manifest["input_sha256"]
    src = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    it = src.active.values; h = list(next(it)); k = h.index("研报文本")
    raw = {i: str(row[k] or "") for i, row in enumerate(it, 2)}
    src.close()
    active, raw_counts, all_events = defaultdict(Counter), defaultdict(Counter), defaultdict(list)
    event_count = 0
    with (args.output_dir / "措辞证据.jsonl").open() as f:
        for line in f:
            e = json.loads(line); text = raw[e["source_row"]]
            assert text[e["start"]:e["end"]] == e["raw_match"]
            assert normalized(e["raw_match"])[0] == e["term"]
            assert text[e["clause_start"]:e["clause_end"]] == e["clause"]
            assert e["category"] in LABELS
            assert e["active"] == (not e["excluded_reason"])
            assert e["strict"] == (e["active"] and not e["context_flags"])
            if e["active"]: active[e["source_row"]][e["category"]] += 1
            raw_counts[e["source_row"]][e["category"]] += 1
            all_events[e["source_row"]].append(e)
            event_count += 1
    docs = json.loads((args.output_dir / "逐篇指标.json").read_text())
    assert [d["source_row"] for d in docs] == list(raw)
    book = openpyxl.load_workbook(args.output_dir / "问题5_措辞强度与极端表达结果.xlsx", read_only=True, data_only=True)
    rows = book.active.values; header = list(next(rows))
    assert book.active.max_row == len(docs) + 1
    for d, values in zip(docs, rows):
        row = dict(zip(header, values)); i = d["source_row"]
        assert row["源Excel行号"] == i
        for cat, label in LABELS.items():
            assert d["counts"][cat] == active[i][cat] == row[label+"次数"]
            assert d["raw_counts"][cat] == raw_counts[i][cat] == row[label+"词面命中数"]
            expected = active[i][cat]*1000/d["han"] if d["han"] else None
            value = row[label+"密度_每千汉字"]
            assert value is None if expected is None else abs(value-expected) < 1e-10
        for name, cats in [("extreme", {"absolute", "hyperbole"}), ("strong", {"absolute", "hyperbole", "certainty", "intensity"})]:
            for prefix, flag in [("", "active"), ("strict_", "strict")]:
                events = [e for e in all_events[i] if e[flag] and e["category"] in cats]
                assert d[prefix+name+"_count"] == len(events)
                assert d[prefix+name+"_sentences"] == len({e["sentence_id"] for e in events})
                assert d[prefix+name+"_sentences"] <= d["sentences"]
        for key, value in row.items():
            if "占比" in key or "比例" in key: assert value is None or 0 <= value <= 1
        for title, prefix, count_key, sent_key in [
            ("极端", "", "extreme_count", "extreme_sentences"),
            ("强势", "", "strong_count", "strong_sentences"),
            ("极端", "strict_", "extreme_count", "extreme_sentences"),
            ("强势", "strict_", "strong_count", "strong_sentences")]:
            density_col = ("保守口径"+title if prefix else title+"措辞")+"密度_每千汉字"
            share_col = ("保守口径"+title if prefix else title+"措辞")+"句占比"
            for col, numerator, denominator in [(density_col, d[prefix+count_key]*1000, d["han"]),
                                                 (share_col, d[prefix+sent_key], d["sentences"])]:
                assert row[col] is None if not denominator else abs(row[col]-numerator/denominator) < 1e-10
        modal_n = d["counts"]["certainty"] + d["counts"]["hedge"]
        assert row["确定性占情态词比例"] is None if not modal_n else abs(
            row["确定性占情态词比例"]-d["counts"]["certainty"]/modal_n) < 1e-10
        assert d["strict_strong_count"] <= d["strong_count"]
        assert d["strict_extreme_count"] <= d["extreme_count"] <= d["strong_count"]
    book.close()
    matches = {}
    if args.replay_dir:
        replay_files = ["措辞证据.jsonl", "逐篇指标.json", "措辞核验样本.csv", "summary.json"]
        if (args.replay_dir / "词典扩展候选_待审核.csv").exists():
            replay_files += ["词典扩展候选_待审核.csv"]
        for name in replay_files:
            matches[name] = sha(args.output_dir / name) == sha(args.replay_dir / name)
            assert matches[name], name
    dump(args.output_dir / "verification.json", {"documents": len(docs), "events": event_count,
         "source_offsets_exact": True, "aggregation_and_excel_values": True,
         "all_sentence_shares_in_range": True, "replay_hash_matches": matches,
         "validation_scope": "软件一致性，不是人工分类准确率", "manifest_sha256": sha(args.output_dir / "manifest.json")})
    print(json.dumps({"documents": len(docs), "events": event_count, "replay": matches}, ensure_ascii=False))


if __name__ == "__main__": main()
