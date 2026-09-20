#!/usr/bin/env python3
"""Deterministic lexical proxies of rhetorical force, not validated author extremism."""
import argparse
import csv
import hashlib
import json
import platform
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
VERSION = "1.0.0"
LABELS = dict(certainty="确定性", absolute="绝对化", hyperbole="夸张修辞", intensity="语气强化",
              hedge="模糊性", magnitude="幅度表达", superlative="最高级与排他表达")
HAN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
BOUNDARY = re.compile(r"<\?>|[。！？!?；;，,：:]+")
SENTENCE_END = re.compile(r"<\?>|[。！？!?]+")
NEGATION = re.compile(r"(?:并不是|并非|并不|不是|没有|未能|未必|未|不|无|没)(?:那么|如此|很|太|十分|非常)?$")
CONDITION = re.compile(r"如果|假如|倘若|一旦|若(?!干)|假设|除非|只要|只有")
ATTRIBUTION = re.compile(r"(?:公司|管理层|董事长|总经理|公告|媒体|市场|业内|客户).{0,8}(?:表示|声称|宣称|称|认为|预计|预期)")
DISCLAIMER = re.compile(r"不构成.{0,12}(?:投资建议|要约)|不保证.{0,15}(?:准确|完整)|(?:本报告|本研报).{0,15}(?:仅供参考|版权|免责声明)|投资者.{0,8}自行.{0,8}(?:承担|判断)")
RATING = re.compile(r"强烈推荐(?:[-－—]?A)?|强力买入|强烈买入")


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def dump(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def normalized(raw):
    """Remove whitespace/soft hyphen, preserving index map to original raw text."""
    positions = [i for i, c in enumerate(raw) if not c.isspace() and c not in "\u00ad\u200b\ufeff"]
    return "".join(raw[i] for i in positions), positions


def spans(text, pattern):
    start = 0
    for m in pattern.finditer(text):
        # ASCII comma within a number is not a clause boundary.
        if m.group() == "," and m.start() and m.end() < len(text) and text[m.start()-1].isdigit() and text[m.end()].isdigit():
            continue
        if m.start() > start:
            yield start, m.start()
        start = m.end()
    if start < len(text):
        yield start, len(text)


class Analyzer:
    def __init__(self, lexicon=HERE / "resources/lexicon.json"):
        self.lexicon_path = Path(lexicon)
        self.spec = json.loads(self.lexicon_path.read_text(encoding="utf-8"))
        self.categories = self.spec["categories"]
        if set(self.categories) != set(LABELS):
            raise ValueError("词表必须包含固定的七类；新增维度需升级程序版本")
        all_terms = [t for terms in self.categories.values() for t in terms]
        if len(all_terms) != len(set(all_terms)) or any(not t for t in all_terms):
            raise ValueError("词条为空或跨类别重复")
        self.lookup = {t: cat for cat, terms in self.categories.items() for t in terms}
        self.pattern = re.compile("|".join(re.escape(t) for t in sorted(all_terms, key=lambda t: (-len(t), t))))
        self.exclusion = re.compile("|".join(self.spec["exclude_patterns"]))

    def analyze(self, raw, names=()):
        raw = "" if raw is None else str(raw)
        text, pos = normalized(raw)
        placeholder = text.lower() in {"无", "暂无", "无内容", "无有效内容", "nan", "none", "null"}
        if not text or placeholder or not HAN.search(text):
            return {"status": "占位文本" if placeholder else "空白或无汉字", "han": 0,
                    "sentences": 0, "counts": {k: 0 for k in LABELS}, "raw_counts": {k: 0 for k in LABELS},
                    "extreme_count": 0, "extreme_sentences": 0, "strong_count": 0, "strong_sentences": 0,
                    "strict_extreme_count": 0, "strict_strong_count": 0, "strict_extreme_sentences": 0,
                    "strict_strong_sentences": 0, "excluded_count": 0, "context_count": 0}, []
        masks = []
        for name in names:
            name = normalized(str(name or ""))[0]
            if len(name) >= 2:
                masks.extend((m.start(), m.end(), "当前公司或券商名称") for m in re.finditer(re.escape(name), text))
        masks.extend((m.start(), m.end(), "固定评级名称") for m in RATING.finditer(text))
        sentence_spans = list(spans(text, SENTENCE_END))
        for a, b in sentence_spans:
            if DISCLAIMER.search(text[a:b]):
                masks.append((a, b, "免责声明句"))
        mask_chars = set(i for a, b, _ in masks for i in range(a, b))
        effective_sentences = [i for i, (a, b) in enumerate(sentence_spans) if any(HAN.fullmatch(text[j]) and j not in mask_chars for j in range(a, b))]
        han = sum(bool(HAN.fullmatch(c)) and i not in mask_chars for i, c in enumerate(text))
        exclusions = [(m.start(), m.end()) for m in self.exclusion.finditer(text)]
        citations = [(m.start(), m.end()) for m in re.finditer(r"《[^《》]*》", text)]
        events = []
        sentence_idx = 0
        for ca, cb in spans(text, BOUNDARY):
            clause = text[ca:cb]
            while sentence_idx + 1 < len(sentence_spans) and ca >= sentence_spans[sentence_idx][1]:
                sentence_idx += 1
            sa, sb = sentence_spans[sentence_idx]
            sent = text[sa:sb]
            matches = list(self.pattern.finditer(clause))
            hedge_context = any(self.lookup[m.group()] == "hedge" for m in matches)
            conditional = bool(CONDITION.search(clause))
            attribution = bool(ATTRIBUTION.search(sent))
            question = sb < len(text) and text[sb] in "？?"
            for m in matches:
                a, b = ca + m.start(), ca + m.end()
                term, category = m.group(), self.lookup[m.group()]
                excluded = next((reason for x, y, reason in masks if a < y and b > x), "")
                if not excluded and any(a < y and b > x for x, y in exclusions):
                    excluded = "金融术语或事实身份搭配"
                if term.startswith("极") and a and text[a-1] in "积南北":
                    excluded = excluded or "跨词误匹配"
                if term == "绝对" and not excluded and not re.match(
                        r"不会|不可能|不能|不是|不应|没有|会|能|是|确定|肯定|正确|错误|安全|可靠|值得|看好|低估|高估", text[b:]):
                    excluded = "绝对的歧义用法_待核"
                before = clause[:m.start()]
                negated = bool(NEGATION.search(before))
                # Phrase-level hedges (不一定/不确定) consume their own negator; no polarity reversal.
                if negated and not excluded:
                    excluded = "前置否定_保守排除"
                flags = []
                if hedge_context and category != "hedge": flags.append("同分句含模糊词")
                if conditional: flags.append("条件表达")
                if attribution: flags.append("疑似转述")
                if question: flags.append("疑问句")
                if any(a < y and b > x for x, y in citations): flags.append("书名号引用")
                events.append({"term": term, "category": category, "start": pos[a], "end": pos[b-1]+1,
                               "raw_match": raw[pos[a]:pos[b-1]+1], "sentence_id": sentence_idx+1,
                               "clause_start": pos[ca], "clause_end": pos[cb-1]+1,
                               "clause": raw[pos[ca]:pos[cb-1]+1], "excluded_reason": excluded,
                               "context_flags": flags, "active": not bool(excluded),
                               "strict": not excluded and not flags})
        counts = Counter(e["category"] for e in events if e["active"])
        raw_counts = Counter(e["category"] for e in events)
        result = {"status": "可计算" if han >= 100 and len(effective_sentences) >= 3 else "短文本_谨慎比较",
                  "han": han, "sentences": len(effective_sentences), "counts": {k: counts[k] for k in LABELS},
                  "raw_counts": {k: raw_counts[k] for k in LABELS},
                  "excluded_count": sum(not e["active"] for e in events),
                  "context_count": sum(e["active"] and bool(e["context_flags"]) for e in events)}
        if not han: result["status"] = "有效正文为空"
        for name, cats in [("extreme", {"absolute", "hyperbole"}), ("strong", {"certainty", "absolute", "hyperbole", "intensity"})]:
            for prefix, flag in [("", "active"), ("strict_", "strict")]:
                hits = [e for e in events if e[flag] and e["category"] in cats]
                result[prefix+name+"_count"] = len(hits)
                result[prefix+name+"_sentences"] = len({e["sentence_id"] for e in hits})
        return result, events


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=HERE.parent / "综合分析结果.xlsx")
    parser.add_argument("--output-dir", type=Path, default=HERE / "output")
    parser.add_argument("--lexicon", type=Path, default=HERE / "resources/lexicon.json")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    analyzer = Analyzer(args.lexicon)
    src = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    source_rows = src.active.iter_rows(values_only=True)
    headings = list(next(source_rows))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "措辞指标"
    metadata = [k for k in ["fordate", "序号", "stkcd", "公司简称", "证券公司"] if k in headings]
    headers = ["源Excel行号"] + metadata + ["计算状态", "有效汉字数", "有效句数",
       "极端措辞次数", "极端措辞密度_每千汉字", "极端措辞句占比",
       "强势措辞次数", "强势措辞密度_每千汉字", "强势措辞句占比",
       "保守口径极端密度_每千汉字", "保守口径强势密度_每千汉字",
       "保守口径极端句占比", "保守口径强势句占比",
       "确定性占情态词比例", "被排除命中数", "语境待核命中数"]
    for label in LABELS.values(): headers += [label+"次数", label+"密度_每千汉字", label+"词面命中数"]
    ws.append(headers)
    totals, status, excluded, top_terms = Counter(), Counter(), Counter(), Counter()
    samples = defaultdict(list)
    docs = []
    evidence_path = args.output_dir / "措辞证据.jsonl"
    with evidence_path.open("w", encoding="utf-8") as evidence:
        for source_row, row in enumerate(source_rows, 2):
            record = dict(zip(headings, row))
            raw = record["研报文本"]
            result, events = analyzer.analyze(raw, [record.get("公司简称"), record.get("证券公司")])
            han, n = result["han"], result["sentences"]
            density = lambda count: count * 1000 / han if han else None
            ratio = lambda count: count / n if n else None
            counts = result["counts"]
            modal_n = counts["certainty"] + counts["hedge"]
            values = [source_row] + [record.get(k) for k in metadata] + [result["status"], han, n,
                result["extreme_count"], density(result["extreme_count"]), ratio(result["extreme_sentences"]),
                result["strong_count"], density(result["strong_count"]), ratio(result["strong_sentences"]),
                density(result["strict_extreme_count"]), density(result["strict_strong_count"]),
                ratio(result["strict_extreme_sentences"]), ratio(result["strict_strong_sentences"]),
                counts["certainty"] / modal_n if modal_n else None,
                result["excluded_count"], result["context_count"]]
            for cat in LABELS: values += [counts[cat], density(counts[cat]), result["raw_counts"][cat]]
            ws.append(values)
            docs.append({"source_row": source_row, **result})
            totals.update(counts)
            status[result["status"]] += 1
            for e in events:
                e = {"source_row": source_row, **e}
                evidence.write(json.dumps(e, ensure_ascii=False, sort_keys=True) + "\n")
                if e["excluded_reason"]: excluded[e["excluded_reason"]] += 1
                if e["active"]: top_terms[e["term"]] += 1
                group = "排除_"+e["excluded_reason"] if not e["active"] else ("语境待核" if e["context_flags"] else e["category"])
                key = hashlib.sha256((str(source_row)+":"+str(e["start"])).encode()).hexdigest()
                samples[group].append((key, {"抽样层": group, "源Excel行号": source_row, "原文起点": e["start"],
                    "原文终点": e["end"], "词条": e["term"], "建议类别": LABELS[e["category"]],
                    "规则计入": e["active"], "排除理由": e["excluded_reason"], "语境标记": "；".join(e["context_flags"]),
                    "原文上下文": str(raw)[max(0, e["start"]-60):e["end"]+80],
                    "人工是否有效": "", "人工类别": "", "标注人": "", "备注": ""}))
            # Randomly ordered no-hit sentences support recall checks (not only successful matches).
            raw_text = str(raw or "")
            for a, b in spans(raw_text, SENTENCE_END):
                if len(HAN.findall(raw_text[a:b])) < 10: continue
                if any(e["start"] < b and e["end"] > a for e in events): continue
                key = hashlib.sha256((str(source_row)+":"+str(a)).encode()).hexdigest()
                samples["无词面命中句"].append((key, {"抽样层": "无词面命中句", "源Excel行号": source_row,
                    "原文起点": a, "原文终点": b, "词条": "", "建议类别": "", "规则计入": False,
                    "排除理由": "", "语境标记": "", "原文上下文": raw_text[a:b],
                    "人工是否有效": "", "人工类别": "", "标注人": "", "备注": ""}))
            # Bound sampling memory without changing deterministic smallest-hash sample.
            if source_row % 1000 == 0:
                for key in samples: samples[key] = sorted(samples[key], key=lambda x: x[0])[:40]
    src.close()
    sample_rows = [v for group in sorted(samples) for _, v in sorted(samples[group], key=lambda x: x[0])[:30]]
    if sample_rows: write_csv(args.output_dir / "措辞核验样本.csv", sample_rows, list(sample_rows[0]))
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 56
    for col, title in enumerate(headers, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 23 if "密度" in title else (22 if title == "计算状态" else 17)
        cell = ws.cell(1, col)
        cell.fill = PatternFill("solid", fgColor="29445C")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            c = cells[0]
            if "占比" in title or "比例" in title: c.number_format = "0.00%"
            elif "密度" in title: c.number_format = "0.000"
            elif title == "fordate": c.number_format = "yyyy-mm-dd"
    out = args.output_dir / "问题5_措辞强度与极端表达结果.xlsx"
    wb.save(out)
    dump(args.output_dir / "逐篇指标.json", docs)
    summary = {"documents": len(docs), "status": dict(status), "category_counts": dict(totals),
        "excluded": dict(excluded), "top_terms": dict(top_terms.most_common(30)), "review_samples": len(sample_rows),
        "reports_with_extreme": sum(d["extreme_count"] > 0 for d in docs),
        "reports_with_strict_extreme": sum(d["strict_extreme_count"] > 0 for d in docs),
        "validation": "完成程序一致性验证后仍需人工测量效度验证；没有已知准确率"}
    dump(args.output_dir / "summary.json", summary)
    manifest = {"version": VERSION, "python": platform.python_version(), "openpyxl": openpyxl.__version__,
        "input_sha256": sha(args.input), "script_sha256": sha(__file__), "lexicon_sha256": sha(args.lexicon),
        "evidence_sha256": sha(evidence_path), "metrics_sha256": sha(args.output_dir / "逐篇指标.json"),
        "sampling": "SHA256(source_row:start)最小30个/层；分层非代表总体样本", "offsets": "Python Unicode codepoint, 0-based half-open"}
    dump(args.output_dir / "manifest.json", manifest)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
