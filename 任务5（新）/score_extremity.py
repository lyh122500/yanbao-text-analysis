#!/usr/bin/env python3
"""任务5（新）：研报语言偏激与极端程度计分。

词表：resources/lexicon.json（5 维度 361 词，合并自任务5词表.xlsx 与 LM 汇总表 15–16 行）
消歧：resources/disambiguation.json（短词规则，依据语料实测分布，见 README）

测度（沿用词表自带的《使用说明》）：
  1. 比例法（主结果）：某维度词出现次数 ÷ 有效汉字数 × 1000
  2. 二元指标（监管视角）：是否含「绝对化/最高级」用语
  3. 参照词比率：确定性占比 = 确定 / (确定 + 模糊)；程度加强净额 = (高程度 − 弱化)
  4. 偏激指数：5 个维度得分的语料百分位等权平均

单一口径：只排除可确定的误匹配（公司名/券商名、免责声明句、评级档位名、
短词消歧、前置否定），其余全部计入。

输出：output/任务5新_偏激极端程度结果.xlsx、措辞证据.jsonl、逐篇指标.json、
      summary.json、manifest.json
"""
import argparse
import bisect
import hashlib
import json
import platform
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
VERSION = "1.0.0"
RES = HERE / "resources"
DIMS = [("certainty", "确定性强度"), ("absolute", "绝对化/最高级"),
        ("emotion", "情绪极端度"), ("hyperbole", "夸张修辞度"),
        ("intensity", "程度加强度")]

HAN = re.compile(r"[㐀-䶿一-鿿]")
SENTENCE_END = re.compile(r"<\?>|[。！？!?]+")
BOUNDARY = re.compile(r"<\?>|[。！？!?；;，,：:]+")
NEGATION = re.compile(r"(?:并不是|并非|并不|不是|没有|未能|未必|未|不|无|没)(?:那么|如此|很|太|十分|非常)?$")
DISCLAIMER = re.compile(r"不构成.{0,12}(?:投资建议|要约)|不保证.{0,15}(?:准确|完整)|(?:本报告|本研报).{0,15}(?:仅供参考|版权|免责声明)")
# 评级档位名整体掩码：词表里的「强烈」「谨慎」在评级名中是档位而非措辞。
# 实测「强烈」2,595 次中 1,969 次（75.9%）、「谨慎」2,407 次中 1,667 次（69.3%）
# 的后续即评级档位；其余词受影响均在 4% 以下。
RATING = re.compile(r"(?:强烈|强力|谨慎|审慎|积极|重点)(?:推荐|增持|买入|卖出|减持|持有|中性|回避|观望)"
                    r"|(?:买入|增持|减持|卖出|持有|中性|回避|观望|推荐)(?:评级|投资评级)")
# 情绪极端度的极性拆分
POS_SUB, NEG_SUB = "极端积极", "极端消极"
# 依词表《使用说明》第 4 条：监管视角的二元指标只看绝对化/最高级
BINARY_DIM = "absolute"


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def dump(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                    encoding="utf-8")


def normalized(raw):
    """去空白/软连字符/零宽字符，保留到原文的位置映射。"""
    pos = [i for i, c in enumerate(raw) if not c.isspace() and c not in "­​﻿"]
    return "".join(raw[i] for i in pos), pos


def spans(text, pattern):
    start = 0
    for m in pattern.finditer(text):
        if m.group() == "," and m.start() and m.end() < len(text) \
                and text[m.start()-1].isdigit() and text[m.end()].isdigit():
            continue                      # 千分位逗号不是分句边界
        if m.start() > start:
            yield start, m.start()
        start = m.end()
    if start < len(text):
        yield start, len(text)


class Scorer:
    def __init__(self, lexicon=RES / "lexicon.json", disambig=RES / "disambiguation.json"):
        self.spec = json.loads(Path(lexicon).read_text(encoding="utf-8"))
        self.dims = {k: v["words"] for k, v in self.spec["dimensions"].items()}
        self.lookup = {w: k for k, words in self.dims.items() for w in words}
        self.disambig = {k: v for k, v in
                         json.loads(Path(disambig).read_text(encoding="utf-8")).items()
                         if not k.startswith("_")}
        active = [w for w in self.lookup if self.disambig.get(w, {}).get("action") != "drop"]
        self.pattern = re.compile("|".join(
            re.escape(t) for t in sorted(active, key=lambda t: (-len(t), t))))

    def _disambiguate(self, term, norm_text, start, end):
        """返回 (是否计入, 排除理由)。规则依据语料实测后续/前置字分布。"""
        r = self.disambig.get(term)
        if not r:
            return True, ""
        if start and norm_text[start-1] in r.get("exclude_before", []):
            return False, "短词消歧_前置"
        nxt = norm_text[end:end+1]
        if r["action"] == "exclude_after" and nxt in r.get("values", []):
            return False, "短词消歧_后置"
        if r["action"] == "require_after" and nxt not in r.get("values", []):
            return False, "短词消歧_非目标搭配"
        return True, ""

    def score(self, raw, names=()):
        raw = "" if raw is None else str(raw)
        text, pos = normalized(raw)
        placeholder = text.lower() in {"无", "暂无", "无内容", "无有效内容", "nan", "none", "null"}
        blank = {"status": "占位文本" if placeholder else "空白或无汉字", "han": 0, "sentences": 0,
                 "counts": {k: 0 for k, _ in DIMS}, "ref_counts": {k: 0 for k, _ in DIMS},
                 "raw_counts": {k: 0 for k, _ in DIMS},
                 "sub_counts": {}, "excluded": 0, "binary": 0}
        if not text or placeholder or not HAN.search(text):
            return blank, []

        masks = []
        for nm in names:                                   # 公司简称/券商名
            nm = normalized(str(nm or ""))[0]
            if len(nm) >= 2:
                masks.extend((m.start(), m.end(), "公司或券商名称")
                             for m in re.finditer(re.escape(nm), text))
        rating_spans = [(m.start(), m.end()) for m in RATING.finditer(text)]
        sentences = list(spans(text, SENTENCE_END))
        for a, b in sentences:                             # 免责声明整句
            if DISCLAIMER.search(text[a:b]):
                masks.append((a, b, "免责声明句"))
        mask_chars = {i for a, b, _ in masks for i in range(a, b)}
        han = sum(bool(HAN.fullmatch(c)) and i not in mask_chars for i, c in enumerate(text))

        events, sidx = [], 0
        for ca, cb in spans(text, BOUNDARY):
            clause = text[ca:cb]
            while sidx + 1 < len(sentences) and ca >= sentences[sidx][1]:
                sidx += 1
            matches = list(self.pattern.finditer(clause))
            for m in matches:
                a, b = ca + m.start(), ca + m.end()
                term, cat = m.group(), self.lookup[m.group()]
                excluded = next((r for x, y, r in masks if a < y and b > x), "")
                # 评级档位名只屏蔽词命中，不减少有效汉字数——否则分母会随词表变化，
                # 破坏「有效汉字数是文本属性」这一前提与跨版本可比性。
                # 判据是命中区间与评级名区间相交（而非邻近），避免殃及旁边的独立措辞。
                if not excluded and any(a < y and b > x for x, y in rating_spans):
                    excluded = "评级档位名"
                if not excluded:
                    ok, why = self._disambiguate(term, text, a, b)
                    if not ok:
                        excluded = why
                if not excluded and NEGATION.search(clause[:m.start()]):
                    excluded = "前置否定"
                events.append({
                    "term": term, "category": cat,
                    "subclass": self.dims[cat][term]["subclass"],
                    "role": self.dims[cat][term]["role"],
                    "start": pos[a], "end": pos[b-1] + 1, "raw_match": raw[pos[a]:pos[b-1]+1],
                    "sentence_id": sidx + 1, "clause": raw[pos[ca]:pos[cb-1]+1],
                    "excluded_reason": excluded, "active": not excluded})

        # 偏激词进分子；参照词（模糊词/弱化词）是分母与对照，不能混入维度计数。
        # 词表的《使用说明》把 role 定义为「偏激词=计入分子，参照词=用于相对指标分母或对照」。
        counts = Counter(e["category"] for e in events if e["active"] and e["role"] == "偏激词")
        ref_counts = Counter(e["category"] for e in events if e["active"] and e["role"] == "参照词")
        raw_counts = Counter(e["category"] for e in events)
        subs = Counter()
        for e in events:
            if e["active"]:
                subs[(e["category"], e["subclass"])] += 1
        result = {
            "status": "可计算" if han >= 100 and len(sentences) >= 3 else "短文本_谨慎比较",
            "han": han, "sentences": len(sentences),
            "counts": {k: counts[k] for k, _ in DIMS},
            "ref_counts": {k: ref_counts[k] for k, _ in DIMS},
            "raw_counts": {k: raw_counts[k] for k, _ in DIMS},
            "sub_counts": {f"{k}|{s}": v for (k, s), v in subs.items()},
            "excluded": sum(not e["active"] for e in events),
            "binary": 1 if counts[BINARY_DIM] > 0 else 0,
        }
        return result, events


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", type=Path, default=HERE.parent / "综合分析结果.xlsx")
    p.add_argument("--output-dir", type=Path, default=HERE / "output")
    p.add_argument("--lexicon", type=Path, default=RES / "lexicon.json")
    p.add_argument("--disambig", type=Path, default=RES / "disambiguation.json")
    args = p.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    scorer = Scorer(args.lexicon, args.disambig)

    src = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    rows = src.active.iter_rows(values_only=True)
    headings = list(next(rows))
    meta = [k for k in ["fordate", "序号", "stkcd", "公司简称", "证券公司"] if k in headings]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "偏激指标"
    headers = ["源Excel行号"] + meta + ["计算状态", "有效汉字数", "有效句数"]
    for _, label in DIMS:
        headers += [f"{label}_次数", f"{label}_密度", f"{label}_句占比"]
    headers += ["情绪极端度_积极密度", "情绪极端度_消极密度",
                "模糊词_次数", "模糊词_密度", "弱化词_次数", "弱化词_密度",
                "确定性占比", "程度加强净额", "含绝对化或最高级", "偏激指数"]
    ws.append(headers)

    docs, status, totals = [], Counter(), Counter()
    ref_totals = Counter()
    sheet_rows, dim_dens = [], []          # 缓冲：偏激指数需全样本百分位，最后统一写入
    evidence_path = args.output_dir / "措辞证据.jsonl"
    with evidence_path.open("w", encoding="utf-8") as ev:
        for row_i, row in enumerate(rows, 2):
            rec = dict(zip(headings, row))
            raw = rec.get("研报文本")
            res, events = scorer.score(raw, [rec.get("公司简称"), rec.get("证券公司")])
            h = res["han"]
            dens = lambda n: n * 1000 / h if h else None
            c = res["counts"]
            n_sent = res["sentences"]
            sub = res["sub_counts"]
            cert = c["certainty"]
            hedge_n = res["ref_counts"]["certainty"]      # 模糊词（参照词）
            weak_n = res["ref_counts"]["intensity"]       # 弱化词（参照词）
            modal = cert + hedge_n
            # 句占比按句去重：同句多次命中只算一次
            per_dim_sent = defaultdict(set)
            for e in events:
                if e["active"] and e["role"] == "偏激词":
                    per_dim_sent[e["category"]].add(e["sentence_id"])
            vals = [row_i] + [rec.get(k) for k in meta] + [res["status"], h, n_sent]
            for key, label in DIMS:
                vals += [c[key], dens(c[key]),
                         (len(per_dim_sent[key]) / n_sent if n_sent else None)]
            vals += [dens(sub.get("emotion|极端积极", 0)), dens(sub.get("emotion|极端消极", 0)),
                     hedge_n, dens(hedge_n), weak_n, dens(weak_n),
                     (cert / modal if modal else None),
                     (dens(c["intensity"]) - dens(weak_n) if h else None),
                     res["binary"]]
            sheet_rows.append(vals)
            # 偏激指数只用「可计算」样本拟合参考分布
            dim_dens.append([dens(c[k]) if res["status"] == "可计算" else None
                             for k, _ in DIMS])
            docs.append({"source_row": row_i, **res,
                         "sentences_by_dim": {k: len(v) for k, v in per_dim_sent.items()}})
            status[res["status"]] += 1
            totals.update(c)
            ref_totals.update(res["ref_counts"])
            for e in events:
                ev.write(json.dumps({"source_row": row_i, **e}, ensure_ascii=False, sort_keys=True) + "\n")
    src.close()

    # 偏激指数：5 个维度密度的语料百分位等权平均（中秩法，同题3）。
    # 权重等权是透明的设计选择，不是估计出的最优权重；各维度分项全部保留。
    ref = [[v for v in col if v is not None] for col in zip(*dim_dens)]
    ref_sorted = [sorted(col) for col in ref]

    def pct(x, col_sorted):
        n = len(col_sorted)
        if not n or x is None:
            return None
        lo = bisect.bisect_left(col_sorted, x)
        hi = bisect.bisect_right(col_sorted, x)
        return 100.0 * (lo + 0.5 * (hi - lo)) / n

    for vals, dims in zip(sheet_rows, dim_dens):
        ps = [pct(x, col) for x, col in zip(dims, ref_sorted)]
        ps = [p for p in ps if p is not None]
        vals.append(sum(ps) / len(ps) if len(ps) == len(DIMS) else None)
        ws.append(vals)

    ws.freeze_panes = "H2"
    ws.auto_filter.ref = ws.dimensions
    for col, title in enumerate(headers, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 22 if "密度" in title else 18
        cell = ws.cell(1, col)
        cell.fill = PatternFill("solid", fgColor="31506B")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cc = cells[0]
            if "占比" in title or "比例" in title:
                cc.number_format = "0.00%"
            elif "密度" in title or "净额" in title:
                cc.number_format = "0.000"
            elif title == "fordate":
                cc.number_format = "yyyy-mm-dd"
    out = args.output_dir / "任务5新_偏激极端程度结果.xlsx"
    wb.save(out)
    dump(args.output_dir / "逐篇指标.json", docs)
    dump(args.output_dir / "summary.json", {
        "documents": len(docs), "status": dict(status),
        "category_counts": dict(totals),
        "reference_word_counts": dict(ref_totals),
        "reports_with_absolute": sum(d["binary"] for d in docs),
        "validation": "程序一致性已验证；无人工效度验证，不报告准确率"})
    dump(args.output_dir / "manifest.json", {
        "version": VERSION, "python": platform.python_version(),
        "openpyxl": openpyxl.__version__,
        "input_sha256": sha(args.input), "script_sha256": sha(Path(__file__)),
        "lexicon_sha256": sha(args.lexicon), "disambig_sha256": sha(args.disambig),
        # 校准文件是 build_lexicon.py 的输入，决定 lexicon.json 的内容，一并锁定
        "calibration_sha256": sha(RES / "确定性词表校准.json"),
        "build_script_sha256": sha(HERE / "build_lexicon.py"),
        "evidence_sha256": sha(evidence_path),
        "metrics_sha256": sha(args.output_dir / "逐篇指标.json"),
        "offsets": "Python Unicode 码位，0 起、左闭右开"})
    print(json.dumps({"documents": len(docs), "status": dict(status),
                      "category_counts": dict(totals),
        "reference_word_counts": dict(ref_totals),
                      "reports_with_absolute": sum(d["binary"] for d in docs)},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
